# -*- coding: utf-8 -*-
"""
输出：
1. fig_q3_moisture_field_a.png：中截面含水率时空热力图；
2. fig_q3_moisture_field_bc.png：全过程径向剖面与末段阈值附近放大图。
"""
from pathlib import Path
import os
import subprocess

import numpy as np
import plotly.graph_objects as go
from openpyxl import load_workbook
from plotly.subplots import make_subplots


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "result3.xlsx"
OUT_HEATMAP = ROOT / "graphics" / "fig_q3_moisture_field_a.png"
OUT_PROFILES = ROOT / "graphics" / "fig_q3_moisture_field_bc.png"

THRESHOLD = 0.15
FONT = "Microsoft YaHei, SimHei, Arial, sans-serif"
GRID = "#D9E2EA"
TEXT = "#23313F"
THRESHOLD_COLOR = "#C84545"
PROFILE_COLORS = {
    6: "#3266A8",
    12: "#4F86C6",
    24: "#3C9D78",
    30: "#8E9B3A",
    36: "#D39A32",
    42: "#E57A30",
    48: "#C4554D",
    54: "#875A9E",
    "end": "#20262E",
}


def apply_windows_kaleido_cleanup_workaround():
    """兼容 Choreographer 在 Windows 上偶发的 taskkill 超时。"""
    if os.name != "nt":
        return
    try:
        import choreographer.browser_async as browser_async
    except ImportError:
        return

    original_kill = browser_async.kill
    if getattr(original_kill, "_cumcm_timeout_workaround", False):
        return

    def kill_without_taskkill_timeout(process):
        try:
            original_kill(process)
        except subprocess.TimeoutExpired:
            pass

    kill_without_taskkill_timeout._cumcm_timeout_workaround = True
    browser_async.kill = kill_without_taskkill_timeout


def read_result():
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    radii = np.asarray(rows[0][1:], dtype=float)
    data = np.asarray(rows[1:], dtype=float)
    times_h = data[:, 0] / 3600.0
    concentration = data[:, 1:]
    wb.close()
    return times_h, radii, concentration


def nearest_profile(times_h, concentration, hour):
    index = int(np.argmin(np.abs(times_h - hour)))
    return concentration[index], float(times_h[index])


def common_layout(fig, width, height, bottom=80):
    fig.update_layout(
        width=width,
        height=height,
        margin=dict(l=90, r=70, t=85, b=bottom),
        font=dict(family=FONT, size=16, color=TEXT),
        paper_bgcolor="white",
        plot_bgcolor="white",
    )
    fig.update_xaxes(
        showline=True,
        linewidth=1.2,
        linecolor="#7B8794",
        ticks="outside",
        showgrid=True,
        gridcolor=GRID,
        gridwidth=1,
        zeroline=False,
        title_font=dict(size=18),
        tickfont=dict(size=15),
    )
    fig.update_yaxes(
        showline=True,
        linewidth=1.2,
        linecolor="#7B8794",
        ticks="outside",
        showgrid=True,
        gridcolor=GRID,
        gridwidth=1,
        zeroline=False,
        title_font=dict(size=18),
        tickfont=dict(size=15),
    )


def plot_heatmap(times_h, radii, concentration):
    # 控制热力图列数以减小论文图片体积，并保留终止时刻。
    stride = max(1, int(np.ceil(len(times_h) / 720)))
    indices = np.arange(0, len(times_h), stride)
    if indices[-1] != len(times_h) - 1:
        indices = np.r_[indices, len(times_h) - 1]
    x = times_h[indices]
    z = concentration[indices].T
    end_h = float(times_h[-1])

    colorscale = [
        [0.00, "#173F5F"],
        [0.18, "#216A8A"],
        [0.38, "#2A9D8F"],
        [0.60, "#8ABF77"],
        [0.80, "#E9C46A"],
        [1.00, "#E76F51"],
    ]
    fig = go.Figure()
    fig.add_trace(
        go.Heatmap(
            x=x,
            y=radii,
            z=z,
            zmin=0.05,
            zmax=2.55,
            colorscale=colorscale,
            colorbar=dict(
                title=dict(text="含水率<br>/(kg/kg)", side="right", font=dict(size=16)),
                thickness=22,
                len=0.88,
                tickfont=dict(size=14),
            ),
            hovertemplate=(
                "时间：%{x:.3f} h<br>半径：%{y:.1f} cm"
                "<br>含水率：%{z:.4f} kg/kg<extra></extra>"
            ),
        )
    )
    fig.add_trace(
        go.Contour(
            x=x,
            y=radii,
            z=z,
            contours=dict(
                start=THRESHOLD,
                end=THRESHOLD,
                size=0.01,
                coloring="none",
                showlabels=False,
            ),
            line=dict(color="white", width=3),
            showscale=False,
            hoverinfo="skip",
            name="0.15 kg/kg 等值线",
            showlegend=True,
        )
    )
    fig.add_vline(
        x=end_h,
        line=dict(color="#FFFFFF", width=2, dash="dash"),
    )
    fig.add_annotation(
        x=end_h,
        y=1.95,
        text=f"严格达标<br>{end_h:.4f} h",
        showarrow=True,
        arrowhead=2,
        ax=-85,
        ay=45,
        bgcolor="rgba(255,255,255,0.92)",
        bordercolor="#66717D",
        borderwidth=1,
        font=dict(size=15, color=TEXT),
    )
    fig.update_layout(
        title=dict(
            text="问题三中截面含水率的时空演化",
            x=0.5,
            xanchor="center",
            font=dict(size=24),
        ),
        legend=dict(
            x=0.015,
            y=0.035,
            bgcolor="rgba(255,255,255,0.86)",
            bordercolor="#AAB3BC",
            borderwidth=1,
            font=dict(size=14),
        ),
    )
    common_layout(fig, width=1500, height=760, bottom=85)
    fig.update_xaxes(title_text="时间 / h", range=[0, end_h + 0.4], dtick=6)
    fig.update_yaxes(title_text="距中心的径向距离 / cm", range=[0, 2.0], dtick=0.25)
    OUT_HEATMAP.parent.mkdir(parents=True, exist_ok=True)
    fig.write_image(OUT_HEATMAP, scale=2)


def add_threshold_band(fig, row, col, y0, y1):
    fig.add_hrect(
        y0=y0,
        y1=THRESHOLD,
        fillcolor="rgba(92, 164, 104, 0.13)",
        line_width=0,
        layer="below",
        row=row,
        col=col,
    )
    fig.add_hline(
        y=THRESHOLD,
        line=dict(color=THRESHOLD_COLOR, width=2.2, dash="dash"),
        row=row,
        col=col,
    )


def plot_profiles(times_h, radii, concentration):
    end_h = float(times_h[-1])
    fig = make_subplots(
        rows=1,
        cols=2,
        column_widths=[0.49, 0.51],
        horizontal_spacing=0.13,
        subplot_titles=[
            "(a) 全过程径向剖面",
            "(b) 末段阈值附近放大",
        ],
    )

    for hour in [6, 12, 24]:
        profile, actual = nearest_profile(times_h, concentration, hour)
        fig.add_trace(
            go.Scatter(
                x=radii,
                y=profile,
                mode="lines+markers",
                line=dict(color=PROFILE_COLORS[hour], width=3),
                marker=dict(size=5),
                name=f"{actual:.0f} h",
                legendgroup="profiles",
                hovertemplate=(
                    f"{actual:.0f} h<br>半径：%{{x:.1f}} cm"
                    "<br>含水率：%{y:.4f} kg/kg<extra></extra>"
                ),
            ),
            row=1,
            col=1,
        )

    for hour in [30, 36, 42, 48, 54]:
        profile, actual = nearest_profile(times_h, concentration, hour)
        fig.add_trace(
            go.Scatter(
                x=radii,
                y=profile,
                mode="lines+markers",
                line=dict(color=PROFILE_COLORS[hour], width=3),
                marker=dict(size=5),
                name=f"{actual:.0f} h",
                legendgroup="profiles",
                hovertemplate=(
                    f"{actual:.0f} h<br>半径：%{{x:.1f}} cm"
                    "<br>含水率：%{y:.4f} kg/kg<extra></extra>"
                ),
            ),
            row=1,
            col=2,
        )

    final_profile = concentration[-1]
    fig.add_trace(
        go.Scatter(
            x=radii,
            y=final_profile,
            mode="lines+markers",
            line=dict(color=PROFILE_COLORS["end"], width=3.6),
            marker=dict(size=5.5, symbol="diamond"),
            name=f"终止 {end_h:.4f} h",
            legendgroup="profiles",
            hovertemplate=(
                f"终止 {end_h:.4f} h<br>半径：%{{x:.1f}} cm"
                "<br>含水率：%{y:.4f} kg/kg<extra></extra>"
            ),
        ),
        row=1,
        col=2,
    )

    add_threshold_band(fig, 1, 1, 0.045, 1.2)
    add_threshold_band(fig, 1, 2, 0.045, 0.215)
    common_layout(fig, width=1600, height=700, bottom=115)
    fig.update_annotations(font=dict(family=FONT, size=20, color=TEXT))
    fig.update_xaxes(title_text="距中心的径向距离 / cm", range=[0, 2.0], dtick=0.25)
    fig.update_yaxes(
        title_text="含水率 / (kg/kg)",
        range=[0.045, 1.15],
        dtick=0.15,
        row=1,
        col=1,
    )
    fig.update_yaxes(
        title_text="含水率 / (kg/kg)",
        range=[0.045, 0.215],
        dtick=0.025,
        row=1,
        col=2,
    )
    # 使用固定位置的直接图例，避免双子图在静态导出时把横向图例压成一列。
    fig.update_traces(showlegend=False)
    labels = [
        ("6 h", PROFILE_COLORS[6]),
        ("12 h", PROFILE_COLORS[12]),
        ("24 h", PROFILE_COLORS[24]),
        ("30 h", PROFILE_COLORS[30]),
        ("36 h", PROFILE_COLORS[36]),
        ("42 h", PROFILE_COLORS[42]),
        ("48 h", PROFILE_COLORS[48]),
        ("54 h", PROFILE_COLORS[54]),
        (f"终止 {end_h:.4f} h", PROFILE_COLORS["end"]),
    ]
    for x_pos, (label, color) in zip(np.linspace(0.035, 0.965, len(labels)), labels):
        fig.add_annotation(
            x=x_pos,
            y=-0.18,
            xref="paper",
            yref="paper",
            text=f"● {label}",
            showarrow=False,
            font=dict(family=FONT, size=14, color=color),
            xanchor="center",
        )
    OUT_PROFILES.parent.mkdir(parents=True, exist_ok=True)
    fig.write_image(OUT_PROFILES, scale=2)


def main():
    times_h, radii, concentration = read_result()
    apply_windows_kaleido_cleanup_workaround()
    plot_heatmap(times_h, radii, concentration)
    plot_profiles(times_h, radii, concentration)
    print(f"已生成：{OUT_HEATMAP}")
    print(f"已生成：{OUT_PROFILES}")
    print(f"终止时刻：{times_h[-1]:.4f} h")


if __name__ == "__main__":
    main()
