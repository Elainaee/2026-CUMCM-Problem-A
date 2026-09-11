"""用 Plotly 绘制 result2.xlsx 中的温度、水分进程图。

先运行 solve_result2.py 生成 result2.xlsx，再运行：python plot_results.py
默认在 ../graphics 文件夹生成三张 PNG，文件名与论文 includegraphics 引用一致。
"""
from pathlib import Path
import argparse
import sys

WORKSPACE_DEPS = Path(__file__).resolve().parents[1] / ".p2deps"
if WORKSPACE_DEPS.exists():
    sys.path.insert(0, str(WORKSPACE_DEPS))

import numpy as np
import openpyxl
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from solve_result2 import detect_critical_time, read_air_data


ROOT = Path(__file__).resolve().parent          # .../CUMCM/Problem II
PROJECT = ROOT.parent                           # .../CUMCM
RESULT2 = PROJECT / "result2.xlsx"


def read_sheet(name, step=30):
    wb = openpyxl.load_workbook(RESULT2, read_only=True, data_only=True)
    ws = wb[name]
    radii = np.asarray([ws.cell(1, col).value for col in range(2, 23)], dtype=float)
    rows = list(ws.iter_rows(min_row=2, values_only=True))[::step]
    wb.close()
    time_h = np.asarray([row[0] for row in rows], dtype=float) / 3600
    values = np.asarray([row[1:22] for row in rows], dtype=float)
    return time_h, radii, values


def field_figure(time_h, radii, values, title, value_label, colorscale, critical_h):
    """左侧显示全过程热力图，右侧显示典型时刻径向剖面。"""
    fig = make_subplots(rows=1, cols=2, column_widths=[0.57, 0.43], horizontal_spacing=0.13,
                        subplot_titles=("全过程时空分布", "典型时刻径向剖面"))
    fig.add_trace(go.Heatmap(
        x=time_h, y=radii, z=values.T, colorscale=colorscale,
        colorbar=dict(x=0.50, len=0.78, thickness=15),
        hovertemplate="时间 %{x:.2f} h<br>距轴线 %{y:.1f} cm<br>数值 %{z:.4f}<extra></extra>",
    ), row=1, col=1)
    colors = ["#4472C4", "#ED7D31", "#70AD47", "#8064A2", "#4BACC6", "#A56A40"]
    for n, hour in enumerate([0.5, 1, 1.5, 2, 2.5, 3]):
        profile = np.asarray([np.interp(hour, time_h, values[:, j]) for j in range(len(radii))])
        fig.add_trace(go.Scatter(x=radii, y=profile, mode="lines", name=f"{hour:g} h",
                                 line=dict(color=colors[n], width=2.4)), row=1, col=2)
    fig.add_vline(x=critical_h, line_color="#C44E52", line_width=2.2, line_dash="dash",
                  annotation_text=f"临界 {critical_h:.3f} h", annotation_font_color="#C44E52",
                  annotation_bgcolor="rgba(255,255,255,0.82)", row=1, col=1)
    fig.update_xaxes(title_text="时间（h）", row=1, col=1)
    fig.update_yaxes(title_text="距轴线距离（cm）", row=1, col=1)
    fig.update_xaxes(title_text="距轴线距离（cm）", row=1, col=2)
    fig.update_yaxes(title_text=value_label, row=1, col=2)
    fig.update_layout(title=title, template="plotly_white", width=1200, height=650,
                      font=dict(family="Arial, PingFang SC, sans-serif", size=15),
                      legend=dict(orientation="h", x=.66, y=-.16, xanchor="center"),
                      margin=dict(l=80, r=55, t=105, b=120))
    return fig


def critical_time_figure(air, critical_time, info):
    t = air[:, 0]
    temperature = air[:, 1]
    sample_dt = float(np.median(np.diff(t)))
    points = int(round(info["window"] / sample_dt)) + 1
    moving_mean = np.full(len(t), np.nan)
    for j in range(points - 1, len(t)):
        moving_mean[j] = temperature[j - points + 1 : j + 1].mean()
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=t / 3600, y=temperature, mode="lines", name="烘房实测温度", line=dict(color="#4472C4", width=1.7)))
    fig.add_trace(go.Scatter(x=t / 3600, y=moving_mean, mode="lines", name="20分钟滑动均温", line=dict(color="#ED7D31", width=2.4)))
    fig.add_hrect(y0=info["plateau"] - info["band"], y1=info["plateau"] + info["band"], fillcolor="#70AD47", opacity=.13, line_width=0, annotation_text="稳态温度带")
    fig.add_vline(x=critical_time / 3600, line_color="#C44E52", line_width=2.5, line_dash="dash", annotation_text=f"临界时刻 {critical_time:.0f} s")
    fig.update_layout(title="预热结束临界时刻判定", xaxis_title="时间（h）", yaxis_title="烘房温度（℃）", template="plotly_white", width=1000, height=620, font=dict(family="Arial, PingFang SC, sans-serif", size=16), legend=dict(orientation="h", y=1.08), margin=dict(l=85, r=55, t=105, b=75))
    return fig


def export_figure(fig, target, stem):
    fig.write_image(str(target / f"{stem}.png"), scale=2)
    print(target / f"{stem}.png")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", default="graphics", help="PNG图片的输出目录（相对 CUMCM 根目录）")
    parser.add_argument("--step", type=int, default=30, help="每隔多少秒取一个点绘图")
    args = parser.parse_args()
    if args.step < 1:
        parser.error("step 必须为正整数")
    target = PROJECT / args.output_dir
    target.mkdir(parents=True, exist_ok=True)
    air = read_air_data(smooth=False)   # 判定图展示实测数据，统计判定亦用原始数据
    critical_time, critical_info = detect_critical_time(air)
    export_figure(critical_time_figure(air, critical_time, critical_info), target, "fig_q2_critical_time")
    for name, title, label, colors in (
        ("温度", "中截面温度时空分布与径向差异", "温度（℃）", "Cividis"),
        ("水分浓度", "中截面含水率时空分布与径向差异", "含水率（kg/kg）", "Viridis"),
    ):
        time_h, radii, values = read_sheet(name, args.step)
        stem = "fig_q2_temperature_field" if name == "温度" else "fig_q2_moisture_field"
        export_figure(field_figure(time_h, radii, values, title, label, colors, critical_time / 3600), target, stem)


if __name__ == "__main__":
    main()
