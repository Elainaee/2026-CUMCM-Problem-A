"""问题二：二维轴对称半长度温度、水分进程并写出 result2.xlsx。

运行：python solve_result2.py
依赖：numpy、scipy、openpyxl。题目附件须位于上级目录的“附件”文件夹。
"""
from pathlib import Path
import argparse
import os
import sys
import time

WORKSPACE_DEPS = Path(__file__).resolve().parents[1] / ".p2deps"
if WORKSPACE_DEPS.exists():
    sys.path.insert(0, str(WORKSPACE_DEPS))

import numpy as np
import openpyxl
from scipy.sparse import coo_matrix, diags
from scipy.sparse.linalg import LinearOperator, cg
from scipy.signal import savgol_filter


ROOT = Path(__file__).resolve().parent          # .../CUMCM/Problem II
PROJECT = ROOT.parent                           # .../CUMCM
ATTACHMENTS = PROJECT / "附件"
sys.path.insert(0, str(PROJECT))
from drying_common import RTOL, ensure_deliverable_writable, save_deliverable

# Savitzky-Golay 保形滤波参数（与论文一致：窗宽 31 点 = 31 min，三阶多项式）
SMOOTH_WIN, SMOOTH_ORDER = 31, 3


def read_air_data(smooth=True):
    """读入附件1；smooth=True 时先做 Savitzky-Golay 保形滤波。"""
    wb = openpyxl.load_workbook(ATTACHMENTS / "附件1.xlsx", read_only=True, data_only=True)
    data = np.asarray(list(wb.active.values)[1:], dtype=float)
    wb.close()
    if smooth:
        data[:, 1] = savgol_filter(data[:, 1], SMOOTH_WIN, SMOOTH_ORDER)
        data[:, 2] = savgol_filter(data[:, 2], SMOOTH_WIN, SMOOTH_ORDER)
    return data


def detect_critical_time(air, tail_start=10800, window=1200, persistence=1800):
    """由平台接近程度和升温斜率共同识别预热结束时刻。

    尾段中位数估计平台温度，MAD估计测温波动。候选时刻之前20分钟
    的均温须进入2倍MAD带，回归斜率须不超过“3倍MAD/窗长”，且
    两个条件此后连续保持30分钟。该判据用于离线分段，避免单点越阈。
    """
    t, temperature = air[:, 0], air[:, 1]
    sample_dt = float(np.median(np.diff(t)))
    tail = temperature[t >= tail_start]
    if len(tail) < 5:
        raise ValueError("附件1尾段数据不足，无法估计稳态")
    plateau = float(np.median(tail))
    sigma = float(1.4826 * np.median(np.abs(tail - plateau)))
    window_points = int(round(window / sample_dt)) + 1
    keep_points = int(round(persistence / sample_dt)) + 1
    mean = np.full(len(t), np.nan)
    slope = np.full(len(t), np.nan)
    for j in range(window_points - 1, len(t)):
        sl = slice(j - window_points + 1, j + 1)
        minutes = (t[sl] - t[sl][0]) / 60
        mean[j] = temperature[sl].mean()
        slope[j] = np.polyfit(minutes, temperature[sl], 1)[0]
    band = 2 * sigma
    slope_limit = 3 * sigma / (window / 60)
    stable = (np.abs(mean - plateau) <= band) & (np.abs(slope) <= slope_limit)
    for j in range(window_points - 1, len(t) - keep_points + 1):
        if np.all(stable[j : j + keep_points]):
            return float(t[j]), {
                "plateau": plateau,
                "sigma": sigma,
                "band": band,
                "slope_limit": slope_limit,
                "window_mean": float(mean[j]),
                "window_slope": float(slope[j]),
                "window": window,
                "persistence": persistence,
            }
    raise RuntimeError("附件1中未识别到满足持续性要求的恒温阶段")


def material(T,C,drying_stage):
    """附录2/3公式对照接口；生产计算由共用核心决定物性阶段。"""
    if np.min(C)<=0:
        raise ValueError("含水率必须为正")
    core=model_core()
    time=core.CRITICAL_TIME if drying_stage else 0.
    rho,cp,k=core.material_properties(C,time)
    return rho*cp,k,core.moisture_diffusivity(C,T,time)


def model_core():
    """复用第三问半长度核心，避免物性、边界和面系数出现两份实现。"""
    import importlib.util
    name="cumcm_problem3_core"
    if name not in sys.modules:
        spec=importlib.util.spec_from_file_location(name, PROJECT/"Problem III/solve_problem3.py")
        module=importlib.util.module_from_spec(spec)
        sys.modules[name]=module
        spec.loader.exec_module(module)
    return sys.modules[name]


def solve(nr=320, nphi=None, nz=40, dt=300., end=10800, export=True, rtol=RTOL,
          material_mode=None, xlsx_path=None):
    """二维轴对称半长度；nz为半长度分段数，dt为BDF最大步长上限。

    前4h环境数据每60s变化，drying_common 的统一规则会把最大步长压到60s，
    因此本问默认 dt=300 与第三、四问同一套时间步策略。
    """
    if nphi is not None:
        raise ValueError("半长度轴对称模型不再使用nphi；请删除周向网格参数")
    if end<=0 or end!=int(end):
        raise ValueError("输出时长必须为正整数秒")
    started=time.time()
    if export:      # 交付件只写 result2.xlsx，占用时在长算前报错
        ensure_deliverable_writable(PROJECT / "result2.xlsx" if xlsx_path is None else Path(xlsx_path))
    core=model_core()
    with core._overrides(MATERIAL_MODE=material_mode or core.MATERIAL_MODE):
        result=core.simulate_bdf(n_radial=nr,n_axial=nz,rtol=rtol,
            max_step=dt,max_duration=end,record_interval=1.,stop_at_threshold=False)
    seconds=result["time_s"][1:].astype(int)
    sampled_T=result["temperature_c"][1:]
    sampled_C=result["moisture_kg_per_kg"][1:]
    for t in range(1800,int(end)+1,1800):
        print(f"{t/3600:.1f} h: 中心温度 {sampled_T[t-1,0]:.6f}，中心含水率 {sampled_C[t-1,0]:.6f}")
    if export:
        write_result(seconds,sampled_T,sampled_C,xlsx_path)
    print("干质量归一化守恒:",result["conservation"])
    print(f"完成，用时 {time.time()-started:.1f} s")
    return seconds,sampled_T,sampled_C


def write_result(seconds, temperature, moisture, output_path=None):
    """保持附件3的两张结果表结构，原子写到 CUMCM 根目录的 result2.xlsx。

    先写同目录临时文件再替换目标；目标被 Excel 等程序占用时删除临时文件并报错，
    不生成带 _pending 的替代结果：交付件始终只有 result2.xlsx。
    """
    target = PROJECT / "result2.xlsx" if output_path is None else Path(output_path)
    wb = openpyxl.load_workbook(ATTACHMENTS / "附件3" / "result2.xlsx")
    for name, values in (("温度", temperature), ("水分浓度", moisture)):
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                cell.value = None
        ws.cell(1, 1, "时间/s；距中心距离/cm")
        for col, radius_cm in enumerate(np.arange(21) / 10, start=2):
            ws.cell(1, col, float(radius_cm))
        for row, t in enumerate(seconds, start=2):
            ws.cell(row, 1, int(t))
            for col, value in enumerate(values[row - 2], start=2):
                ws.cell(row, col, round(float(value), 4)).number_format = "0.0000"
        ws.freeze_panes = "B2"
    return save_deliverable(wb, target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="计算问题二的二维轴对称半长度温度、水分进程")
    parser.add_argument("--nr", type=int, default=320, help="径向分段数；320=20×16，交付表1mm列正好是节点")
    parser.add_argument("--nz", type=int, default=40, help="半长度轴向分段数")
    parser.add_argument("--dt", type=float, default=300, help="BDF最大时间步长上限/s（前4h按共用规范压到60s）")
    parser.add_argument("--end", type=int, default=10800, help="结束时间/s")
    parser.add_argument("--rtol", type=float, default=RTOL)
    parser.add_argument("--material-mode", choices=["specified","staged"], default=None)
    args = parser.parse_args()
    if args.nr < 2 or args.nz < 2 or args.dt <= 0 or args.rtol <= 0:
        parser.error("网格数过小或时间步长不合法")
    solve(nr=args.nr, nz=args.nz, dt=args.dt, end=args.end,
          rtol=args.rtol, material_mode=args.material_mode)
